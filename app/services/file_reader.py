from __future__ import annotations

import csv
from dataclasses import dataclass
from io import BytesIO, StringIO
from pathlib import Path

import openpyxl

from app.exceptions import AppError
from app.services.normalization import build_column_index, is_empty_data_row, normalize_row


@dataclass(frozen=True)
class ParsedDataset:
    rows: list[dict[str, str | int]]
    detected_columns: list[str]


def parse_dataset_file(filename: str, content: bytes) -> ParsedDataset:
    if not content:
        raise AppError(
            400,
            "archivo_vacio",
            "El archivo enviado esta vacio.",
            {"archivo": filename, "sugerencia": "Seleccione un archivo .xlsx o .csv con datos."},
        )

    suffix = Path(filename).suffix.casefold()
    if suffix == ".xlsx":
        return _parse_xlsx(content)
    if suffix == ".csv":
        return _parse_csv(content)

    raise AppError(
        415,
        "formato_invalido",
        "Formato de archivo invalido. Solo se aceptan archivos .xlsx o .csv.",
        {
            "archivo": filename,
            "formatoRecibido": suffix or "sin extension",
            "formatosPermitidos": [".xlsx", ".csv"],
        },
    )


def _parse_xlsx(content: bytes) -> ParsedDataset:
    try:
        workbook = openpyxl.load_workbook(BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise AppError(
            400,
            "excel_invalido",
            "No fue posible leer el archivo Excel.",
            {"sugerencia": "Verifique que el archivo no este corrupto y que sea .xlsx real."},
        ) from exc

    sheet = workbook[workbook.sheetnames[0]]
    rows_iterator = sheet.iter_rows(values_only=True)

    try:
        headers = list(next(rows_iterator))
    except StopIteration as exc:
        raise AppError(
            400,
            "archivo_vacio",
            "El archivo Excel no contiene filas.",
            {"sugerencia": "Incluya una fila de encabezados y al menos una fila de datos."},
        ) from exc

    column_index, detected_columns, missing = build_column_index(headers)
    _raise_missing_columns_if_needed(missing)

    rows: list[dict[str, str | int]] = []
    for source_row, values in enumerate(rows_iterator, start=2):
        row = normalize_row(values, column_index, source_row)
        if not is_empty_data_row(row):
            rows.append(row)
    return ParsedDataset(rows=rows, detected_columns=detected_columns)


def _parse_csv(content: bytes) -> ParsedDataset:
    decoded = _decode_csv(content)
    stream = StringIO(decoded)
    sample = decoded[:4096]

    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel

    reader = csv.reader(stream, dialect)
    try:
        headers = next(reader)
    except StopIteration as exc:
        raise AppError(
            400,
            "archivo_vacio",
            "El archivo CSV no contiene filas.",
            {"sugerencia": "Incluya una fila de encabezados y al menos una fila de datos."},
        ) from exc

    column_index, detected_columns, missing = build_column_index(headers)
    _raise_missing_columns_if_needed(missing)

    rows: list[dict[str, str | int]] = []
    for source_row, values in enumerate(reader, start=2):
        row = normalize_row(values, column_index, source_row)
        if not is_empty_data_row(row):
            rows.append(row)
    return ParsedDataset(rows=rows, detected_columns=detected_columns)


def _decode_csv(content: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "latin-1"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    return content.decode("utf-8", errors="replace")


def _raise_missing_columns_if_needed(missing: list[str]) -> None:
    if missing:
        raise AppError(
            422,
            "columnas_faltantes",
            "Faltan columnas obligatorias: " + ", ".join(missing),
            {
                "columnasFaltantes": missing,
                "sugerencia": "Revise que los encabezados coincidan con la plantilla esperada.",
            },
        )
